import os
import csv
import re
from dataclasses import asdict
from datetime import datetime
from typing import List, Optional, Tuple
from urllib.parse import urljoin, urlparse
import requests
import zipfile
from pathlib import Path
from urllib.parse import urlsplit

from ccgp.config import (
    PROVINCES,
    RE_DATE_YMD_HM,
    RE_MONEY,
    SG_TZ,
    FILTER_EXCLUDE_KEYWORDS,
    REQUEST_TIMEOUT_SEC,
    DOWNLOAD_TIMEOUT_SEC,
)
from ccgp.model import TenderItem
from utils.mylogger import get_logger

#------------------------------通用工具函数-----------------------------------#
def http_get(url: str, session: requests.Session, timeout: int = REQUEST_TIMEOUT_SEC) -> str:
    resp = session.get(url, timeout=timeout)
    resp.raise_for_status()
    resp.encoding = resp.apparent_encoding or "utf-8"
    return resp.text

def norm_list_page_urls(start_url: str, max_pages: int = 30) -> List[str]:
    """
    处理start_url，生成分页列表页 URL 列表。
    """
    u = start_url.strip()
    parsed = urlparse(u)
    path = parsed.path

    base_dir = u.rsplit("/", 1)[0] + "/"

    urls = [] # 分页URL列表
    for i in range(max_pages):
        if i == 0:
            urls.append(urljoin(base_dir, "index.htm"))
        else:
            urls.append(urljoin(base_dir, f"index_{i}.htm"))
    return urls

def parse_pub_datetime(text: str) -> Optional[datetime]:
    if not text:
        return None
    m = RE_DATE_YMD_HM.search(text)
    if not m:
        return None
    y = int(m.group("y"))
    mo = int(m.group("m"))
    d = int(m.group("d"))
    h = int(m.group("h") or 0)
    mi = int(m.group("mi") or 0)
    try:
        return datetime(y, mo, d, h, mi, tzinfo=SG_TZ)
    except Exception:
        return None

def keyword_hit(text: str, keywords: List[str]) -> bool:
    # 将文本分割成句子（简单的句子分割，可根据需要改进）
    sentences = text.replace(';', ' ').replace('。', ' ').split(' ')
    
    for k in keywords:
        for sentence in sentences:
            if k in sentence.strip():
                flag = True
                for exclude in FILTER_EXCLUDE_KEYWORDS: # 过滤干扰词
                    if exclude in sentence: 
                        flag = False
                        break 
                if flag:                     
                    get_logger().debug(f"Keyword hit: '{k}' in sentence: '{sentence.strip()}'")
                    return True
    return False

def clean_text(s: str) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s

def extract_money(text: str) -> str:
    text = text.replace("人民币", "")
    m = RE_MONEY.search(text)
    
    if not m:
        return ""
    
    num_str = m.group("num") or ""
    unit = m.group("unit") or ""
    
    num_str_clean = num_str.replace(",", "")
    try:
        num = float(num_str_clean)
        # 保留两位小数
        formatted_num = f"{num:.2f}"
        
        if not unit: # 默认单位为“万元”
            return f"{formatted_num}万元"
        else:
            return f"{formatted_num}{unit}"
            
    except (ValueError, AttributeError):
        # 如果转换失败，返回原始格式
        if not unit:
            return f"{num_str}万元"
        else:
            return f"{num_str}{unit}".strip()


def guess_location(addr_text: str) -> Tuple[str, str]:
    """
    从地址文本里找省市
    """
    t = addr_text or ""
    #省
    prov = ""
    for p in PROVINCES:
        if p in t:
            prov = p
            break
    city = ""
    #市
    m = re.search(r"([^\s，,。省（(]{2,20}市)", t)
    if m:
        city = m.group(1)
    return prov, city

def write_csv(items: List[TenderItem], out_path: str) -> None:
    cols = [
        "project_name", "pub_time",
        "ai_project_title","requirement_brief", "requirement_desc","deadline",
        "company_name", "contact_name", "contact_phone",
        "province", "city",
        "budget","announcement_url","location_text"
    ]
    directory = os.path.dirname(out_path)
    if directory and not os.path.exists(directory):
        os.makedirs(directory, exist_ok=True)
        
    file_exists = os.path.exists(out_path) and os.path.getsize(out_path) > 0
    
    # 1. 收集所有行数据（含既有数据和新传入的数据）
    all_rows = []
    
    if file_exists:
        try:
            with open(out_path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                all_rows.extend(list(reader))
        except Exception as e:
            get_logger().error(f"Error reading existing CSV: {e}")
            
    # 追加新传入的项目
    for it in items:
        row = asdict(it)
        all_rows.append({k: row.get(k, "") for k in cols})

    # 2. 根据 announcement_url 去重，保留最新（即列表靠后）读到的项
    unique_rows = {}
    for r in all_rows:
        url = r.get("announcement_url", "")
        unique_rows[url] = r
    
    all_rows = list(unique_rows.values())
    
    # 3. 按照 pub_time 降序排序 (越新的越靠前)
    # 假设 pub_time 是格式为 '%Y-%m-%d %H:%M:%S' 或类似字符串
    # 无法解析时间的默认放在最后面
    def sort_key(row):
        t_str = row.get("pub_time", "")
        try:
            # 兼容带有时间或者仅有日期的格式，无法解析的返回极小时间
            return datetime.strptime(t_str[:16], "%Y-%m-%d %H:%M")
        except Exception:
            try:
                return datetime.strptime(t_str[:10], "%Y-%m-%d")
            except Exception:
                return datetime.min

    all_rows.sort(key=sort_key, reverse=True)

    # 4. 全量覆盖写入
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(all_rows)

#------------------------------和附件相关-----------------------------------#
def safe_filename(name: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name or "").strip()
    return name[:120] or "file"

def download_file(
    session: requests.Session,
    url: str,
    out_dir: str,
    filename: Optional[str] = None,
    max_bytes: int = 35 * 1024 * 1024,
    timeout: int = DOWNLOAD_TIMEOUT_SEC,
) -> str:

    Path(out_dir).mkdir(parents=True, exist_ok=True)
    if filename:
        fn = safe_filename(Path(filename).name)
    else:
        fn = safe_filename(Path(urlsplit(url).path).name)  #URL path 
    if not fn or "." not in fn:
        fn = safe_filename(fn) + ".bin"
    path = str(Path(out_dir) / fn)

    if os.path.exists(path) and os.path.getsize(path) > 0:
        return path

    with session.get(url, stream=True, timeout=timeout) as r:
        r.raise_for_status()
        total = 0
        with open(path, "wb") as f:
            for chunk in r.iter_content(chunk_size=8192):
                if not chunk:
                    continue
                total += len(chunk)
                if total > max_bytes:
                    raise ValueError(f"file too large > {max_bytes} bytes: {url}")
                f.write(chunk)
    return path


def extract_text_from_pdf(path: str) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        parts = []
        for p in reader.pages:
            t = p.extract_text() or ""
            if t.strip():
                parts.append(t)
        return "\n".join(parts)
    except Exception as e:
        raise ValueError("Failed to open PDF file: " + str(e))


def extract_text_from_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if (p.text or "").strip()]
        return "\n".join(parts)
    except Exception as e:
        raise ValueError("Failed to open DOCX file: " + str(e))
        


def extract_text_from_xlsx(path: str, max_cells: int = 20000) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        out = []
        cells = 0
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if cells > max_cells:
                    out.append("...(truncated)")
                    return "\n".join(out)
                line = []
                for v in row:
                    if v is None:
                        continue
                    s = str(v).strip()
                    if s:
                        line.append(s)
                if line:
                    out.append(" | ".join(line))
                cells += len(row)
        return "\n".join(out)
    except Exception as e:
        raise ValueError("Failed to open XLSX file: " + str(e))


def extract_text_from_zip(path: str, tmp_dir: str) -> str:
    """
    只解析 zip 内常见可提取格式（pdf/docx/xlsx/txt），避免炸弹：限制文件数/单文件大小
    """
    texts = []
    try:
        Path(tmp_dir).mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path, "r") as z:
            names = z.namelist()[:30]  # 限制最多 30 个文件
            for n in names:
                nl = n.lower()
                if not (nl.endswith(".pdf") or nl.endswith(".docx") or nl.endswith(".xlsx") or nl.endswith(".txt")):
                    continue
                info = z.getinfo(n)
                if info.file_size > 30 * 1024 * 1024:
                    continue
                extracted = z.extract(n, path=tmp_dir)
                texts.append(f"\n### ZIP_ENTRY: {n}\n")
                texts.append(extract_text_from_file(extracted))
        return "\n".join(texts)
    except Exception as e:
        raise ValueError("Failed to open ZIP file: " + str(e))


def extract_text_from_file(path: str) -> str:
    p = path.lower()
    if p.endswith(".pdf"):
        return extract_text_from_pdf(path)
    if p.endswith(".docx"):
        return extract_text_from_docx(path)
    if p.endswith(".xlsx"):
        return extract_text_from_xlsx(path)
    if p.endswith(".txt"):
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""
    if p.endswith(".zip"):
        return extract_text_from_zip(path, tmp_dir=str(Path(path).parent / "_unzipped"))
    # .doc/.xls/.rar 这里先不处理（后续可用转换或额外库）
    return ""
